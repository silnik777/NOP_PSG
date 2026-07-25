"""API tests for Faza II: hydraulics, linepack and export."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

_HYDRAULICS_BODY = {
    "compositionId": "REF-GAS-E",
    "diameter": {"value": 0.5, "unit": "m"},
    "roughness": {"value": 0.012, "unit": "mm"},
    "length": {"value": 50, "unit": "km"},
    "inletPressure": {"value": 5.0, "unit": "MPa"},
    "gasTemperature": {"value": 283.15, "unit": "K"},
    "normalFlow": {"value": 100000, "unit": "Nm3/h"},
}


def test_steady_flow_endpoint(client):
    r = client.post("/api/v1/hydraulics/steady-flow", json=_HYDRAULICS_BODY)
    assert r.status_code == 200, r.text
    data = r.json()
    assert 0 < data["pressureDrop"]["value"] < 5.0
    assert data["flowRegime"] == "Turbulent"
    assert data["frictionFactor"] > 0
    assert data["resultClass"] == "Engineering"


def test_linepack_endpoint(client):
    r = client.post("/api/v1/hydraulics/linepack", json=_HYDRAULICS_BODY)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["linepackMass"]["value"] > 0
    assert data["linepackMass"]["unit"] == "t"
    assert data["linepackNormalVolume"]["value"] > 0


def test_export_hydraulics_xlsx(client):
    r = client.post("/api/v1/export/hydraulics?format=xlsx", json=_HYDRAULICS_BODY)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    banner = ws.cell(row=1, column=1).value
    assert banner.startswith("[ENGINEERING]")
    labels = [ws.cell(row=i, column=1).value for i in range(1, ws.max_row + 1)]
    assert "Pressure drop" in labels
    assert "Friction factor (Darcy)" in labels


def test_export_compression_csv(client):
    body = {
        "compositionId": "REF-GAS-H2-20",
        "massFlowRate": {"value": 10.0, "unit": "kg/s"},
        "inletPressure": {"value": 2.0, "unit": "MPa"},
        "inletTemperature": {"value": 288.15, "unit": "K"},
        "outletPressureTarget": {"value": 5.0, "unit": "MPa"},
        "isentropicEfficiency": 0.80,
    }
    r = client.post("/api/v1/export/compression?format=csv", json=body)
    assert r.status_code == 200, r.text
    assert "text/csv" in r.headers["content-type"]
    text = r.content.decode("utf-8")
    assert "[ENGINEERING]" in text
    assert "Required shaft power" in text


def test_export_rejects_bad_format(client):
    r = client.post("/api/v1/export/hydraulics?format=pdf", json=_HYDRAULICS_BODY)
    assert r.status_code == 422
